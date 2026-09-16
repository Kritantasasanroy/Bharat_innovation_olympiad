from openpyxl import load_workbook
from pathlib import Path

file = Path(r"C:\KSR\Lemon Ideas\Bharat_innovation_olympiad\STUDENT_PORTAL_CONTENT_CATALOG.xlsx")
wb = load_workbook(file, data_only=True)
ws = wb['Student Content']
headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
file_idx = headers.index('Code File')
loc_idx = headers.index('Line(s) / Searchable Location')
content_idx = headers.index('Current Content')
for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    if row[file_idx] == 'frontend/src/components/TooSmallForExam.tsx':
        print(f"Row {i} ({row[loc_idx]}): {row[content_idx]!r}")
