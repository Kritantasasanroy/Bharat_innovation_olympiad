from openpyxl import load_workbook
from pathlib import Path

p = Path(r"C:\KSR\Lemon Ideas\Bharat_innovation_olympiad\STUDENT_PORTAL_CONTENT_CATALOG (1).xlsx")
wb = load_workbook(p, data_only=True)
ws = wb["Student Content"]
headers = [c.value for c in ws[1]]
new_idx = headers.index("New Content")
cur_idx = headers.index("Current Content")
total = 0
new_filled = 0
cur_filled = 0
changed = 0
for row in ws.iter_rows(min_row=2, values_only=True):
    total += 1
    if row[new_idx]:
        new_filled += 1
        if new_filled <= 5:
            print("New:", repr(row[new_idx]))
    if row[cur_idx]:
        cur_filled += 1
    if row[new_idx] and row[cur_idx] and row[new_idx].strip() != row[cur_idx].strip():
        changed += 1
        if changed <= 5:
            print("Changed row:", total + 1, "\n  Current:", repr(row[cur_idx]), "\n  New:", repr(row[new_idx]))
print(f"Total rows: {total}, Current Content filled: {cur_filled}, New Content filled: {new_filled}, Changed: {changed}")
