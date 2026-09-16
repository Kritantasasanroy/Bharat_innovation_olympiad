from openpyxl import load_workbook
from pathlib import Path
import json

ROOT = Path(r"C:\KSR\Lemon Ideas\Bharat_innovation_olympiad")
wb = load_workbook(ROOT / "STUDENT_PORTAL_CONTENT_CATALOG (1).xlsx", data_only=True)
ws = wb["Student Content"]
headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]

rows = []
for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    d = {h: v for h, v in zip(headers, row)}
    d["row"] = i
    rows.append(d)

out = ROOT / "catalog.json"
out.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"Wrote {len(rows)} rows to {out}")
