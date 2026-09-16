from openpyxl import load_workbook
from pathlib import Path

p = Path(r"C:\KSR\Lemon Ideas\Bharat_innovation_olympiad\STUDENT_PORTAL_CONTENT_CATALOG (1).xlsx")
wb = load_workbook(p, data_only=True)
ws = wb["Student Content"]
headers = [c.value for c in ws[1]]
new_idx = headers.index("New Content")
cur_idx = headers.index("Current Content")
updates = []
for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    new = str(row[new_idx] or '').strip()
    cur = str(row[cur_idx] or '').strip()
    if new and new.lower() != 'ok':
        updates.append((i, row[headers.index('Portal')], row[headers.index('Area / Page / Feature')], row[headers.index('Code File')], row[headers.index('Line(s) / Searchable Location')], cur, new))
print(f"Rows with actual New Content: {len(updates)} / {ws.max_row - 1}")
for u in updates[:50]:
    print(f"Row {u[0]}: {u[3]} {u[4]}")
    print(f"  Current: {u[5]!r}")
    print(f"  New:     {u[6]!r}")
    print()
if len(updates) > 50:
    print(f"... plus {len(updates)-50} more")
