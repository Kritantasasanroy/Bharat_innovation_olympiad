from openpyxl import load_workbook

wb = load_workbook(
    r"C:\KSR\Lemon Ideas\Bharat_innovation_olympiad\STUDENT_PORTAL_CONTENT_CATALOG.xlsx",
    data_only=True,
)
ws = wb["Student Content"]
headers = [c.value for c in ws[1]]
new_idx = headers.index("New Content")
total = 0
filled = 0
for row in ws.iter_rows(min_row=2, values_only=True):
    total += 1
    if row[new_idx]:
        filled += 1
        if filled <= 10:
            print(repr(row[new_idx]))
print(f"Total rows: {total}, New Content filled: {filled}")
